"""Leitura de documentos do chat com os loaders do LangChain.

PDF → PyPDFLoader, Word → Docx2txtLoader, CSV → CSVLoader, texto → TextLoader.
Excel usa um loader próprio no mesmo contrato (BaseLoader) sobre o openpyxl: o
UnstructuredExcelLoader exige o pacote `unstructured`, pesado demais para o
instalador. Tudo roda neste PC; nada é enviado para fora.

Os imports do LangChain ficam dentro das funções: o backend sobe sem pagar o
custo deles até alguém anexar um arquivo.
"""

from __future__ import annotations

from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

SUPPORTED = {
    ".pdf": "pdf",
    ".docx": "word",
    ".xlsx": "excel",
    ".xlsm": "excel",
    ".csv": "csv",
    ".txt": "text",
    ".md": "text",
}
CHUNK_SIZE = 1_200
CHUNK_OVERLAP = 150
MAX_ROWS_PER_SHEET = 5_000


class DocumentError(ValueError):
    """Código curto (unsupported, empty, unreadable) que a interface traduz."""


@dataclass(frozen=True)
class Piece:
    text: str
    page: int | None


def document_kind(filename: str) -> str:
    kind = SUPPORTED.get(Path(filename).suffix.lower())
    if kind is None:
        raise DocumentError("unsupported")
    return kind


def read_document(path: Path, filename: str) -> tuple[list[Piece], int]:
    """Lê e fatia o arquivo. Devolve os trechos e o número de páginas/abas."""
    kind = document_kind(filename)
    try:
        documents = _loader(kind, path).load()
    except DocumentError:
        raise
    except Exception as exc:  # PDF protegido, Word corrompido, planilha inválida…
        raise DocumentError("unreadable") from exc
    documents = [doc for doc in documents if doc.page_content and doc.page_content.strip()]
    if not documents:
        raise DocumentError("empty")

    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
    pieces = [
        Piece(text=chunk.page_content.strip(), page=_page(chunk.metadata, kind))
        for chunk in splitter.split_documents(documents)
        if chunk.page_content.strip()
    ]
    pages = len({piece.page for piece in pieces if piece.page is not None}) or 1
    return pieces, pages


def _page(metadata: dict, kind: str) -> int | None:
    if kind == "pdf" and isinstance(metadata.get("page"), int):
        return int(metadata["page"]) + 1  # PyPDFLoader conta a partir de 0
    if kind == "excel" and isinstance(metadata.get("sheet_index"), int):
        return int(metadata["sheet_index"]) + 1
    if kind == "csv" and isinstance(metadata.get("row"), int):
        return None
    return None


def _loader(kind: str, path: Path):
    if kind == "pdf":
        from langchain_community.document_loaders.pdf import PyPDFLoader

        return PyPDFLoader(str(path))
    if kind == "word":
        from langchain_community.document_loaders.word_document import Docx2txtLoader

        return Docx2txtLoader(str(path))
    if kind == "csv":
        from langchain_community.document_loaders.csv_loader import CSVLoader

        return CSVLoader(str(path), autodetect_encoding=True)
    if kind == "text":
        from langchain_community.document_loaders.text import TextLoader

        return TextLoader(str(path), autodetect_encoding=True)
    if kind == "excel":
        return _excel_loader(path)
    raise DocumentError("unsupported")


def _excel_loader(path: Path):
    from langchain_core.document_loaders import BaseLoader
    from langchain_core.documents import Document

    class ExcelSheetsLoader(BaseLoader):
        """Uma aba por documento; cada linha vira "Coluna: valor; Coluna: valor"."""

        def __init__(self, file_path: Path) -> None:
            self.file_path = file_path

        def lazy_load(self) -> Iterator[Document]:
            from openpyxl import load_workbook

            book = load_workbook(self.file_path, read_only=True, data_only=True)
            try:
                for index, sheet in enumerate(book.worksheets):
                    lines = _sheet_lines(sheet.iter_rows(values_only=True))
                    if lines:
                        yield Document(
                            page_content=f"Aba: {sheet.title}\n" + "\n".join(lines),
                            metadata={"source": str(self.file_path), "sheet": sheet.title, "sheet_index": index},
                        )
            finally:
                book.close()

    return ExcelSheetsLoader(path)


def _sheet_lines(rows) -> list[str]:
    header: list[str] | None = None
    lines: list[str] = []
    for row in rows:
        cells = ["" if value is None else str(value).strip() for value in row]
        if not any(cells):
            continue
        if header is None:
            header = [cell or f"Coluna {position + 1}" for position, cell in enumerate(cells)]
            continue
        pairs = [f"{header[i] if i < len(header) else f'Coluna {i + 1}'}: {cell}" for i, cell in enumerate(cells) if cell]
        lines.append("; ".join(pairs))
        if len(lines) >= MAX_ROWS_PER_SHEET:
            lines.append("…")
            break
    if header is not None and not lines:
        lines.append("; ".join(cell for cell in header if cell))
    return lines
