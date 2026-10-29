"""Exportação da conversa em PDF e Word com a cara do DroidNote.

A ata chega em Markdown (títulos, listas, marcas `mm:ss`); os dois formatos a desenham com a
mesma estrutura: título laranja, metadados discretos, seções com régua, marcadores e as marcas
de tempo em fonte mono cinza. A transcrição vem depois, com falante em negrito.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from app.core.i18n import ui_message
from app.domain.models import Person, Session, Summary, TranscriptSegment

ACCENT = "#FF4400"
INK = "#0C0C0E"
SOFT = "#34343A"
MUTED = "#6B6B73"
LINE = "#D9D9DE"
_STAMP = re.compile(r"`(\d{1,2}:\d{2}(?::\d{2})?)`|\((\d{1,2}:\d{2}(?::\d{2})?)\)")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


@dataclass(frozen=True)
class Block:
    kind: str  # title, h2, h3, p, li
    text: str
    level: int = 0


def parse_markdown(markdown: str) -> list[Block]:
    """Markdown da ata → blocos simples. Linhas soltas viram parágrafos; recuo vira nível da lista."""
    blocks: list[Block] = []
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            blocks.append(Block("p", " ".join(paragraph).strip()))
            paragraph.clear()

    for raw in markdown.splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("# "):
            flush()
            blocks.append(Block("title", stripped[2:].strip()))
        elif stripped.startswith("## "):
            flush()
            blocks.append(Block("h2", stripped[3:].strip()))
        elif stripped.startswith("### ") or stripped.startswith("#### "):
            flush()
            blocks.append(Block("h3", stripped.lstrip("#").strip()))
        elif re.match(r"^(?:[-*•]|\d+[.)])\s+", stripped):
            flush()
            indent = (len(line) - len(line.lstrip(" \t"))) // 2
            text = re.sub(r"^(?:[-*•]|\d+[.)])\s+", "", stripped)
            text = re.sub(r"^\[[ xX]\]\s*", "", text)
            blocks.append(Block("li", text, indent))
        else:
            paragraph.append(stripped)
    flush()
    return blocks


def _clock(ms: int) -> str:
    seconds = max(ms, 0) // 1000
    minutes, sec = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{sec:02d}"
    return f"{minutes:02d}:{sec:02d}"


def _duration(session: Session, segments: list[TranscriptSegment]) -> str:
    if session.ended_at:
        total = int((session.ended_at - session.started_at).total_seconds())
    elif segments:
        total = segments[-1].end_ms // 1000
    else:
        total = 0
    return _clock(total * 1000)


def _speaker_name(segment: TranscriptSegment, people: dict[str, str]) -> str:
    if segment.speaker_id and segment.speaker_id in people:
        return people[segment.speaker_id]
    if segment.speaker_cluster:
        from app.application.summarize import voice_label

        return voice_label(segment.speaker_cluster, segment.language)
    return ""


def _participants(session_people: list[Person]) -> str:
    return ", ".join(item.name for item in session_people)


def _summary_markdown(summary: Summary | None, locale: str | None) -> str:
    """A ata em Markdown; sem o texto salvo, monta a partir dos campos."""
    if summary is None:
        return ""
    if summary.notes_markdown.strip():
        return summary.notes_markdown
    lines: list[str] = []
    if summary.overview:
        lines += [f"## {ui_message(locale, 'export_summary')}", summary.overview, ""]
    if summary.highlights:
        lines += [f"## {ui_message(locale, 'export_highlights')}"] + [f"- {item}" for item in summary.highlights] + [""]
    if summary.decisions:
        lines += [f"## {ui_message(locale, 'export_decisions')}"] + [f"- {item}" for item in summary.decisions] + [""]
    if summary.action_items:
        lines.append(f"## {ui_message(locale, 'export_owners')}")
        for item in summary.action_items:
            extra = " — ".join(part for part in [item.owner, item.due] if part)
            lines.append(f"- {item.text}" + (f" ({extra})" if extra else ""))
    return "\n".join(lines)


# ----------------------------------------------------------------------- PDF


def export_pdf(
    session: Session,
    segments: list[TranscriptSegment],
    summary: Summary | None,
    people: list[Person],
    locale: str | None = None,
    participants: list[Person] | None = None,
) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

    body_font, bold_font, mono_font = _register_pdf_fonts()
    names = {item.id: item.name for item in people}
    st = {
        "title": ParagraphStyle("DnTitle", fontName=bold_font, fontSize=20, leading=24, textColor=colors.HexColor(ACCENT), spaceAfter=4),
        "meta": ParagraphStyle("DnMeta", fontName=body_font, fontSize=9, leading=12, textColor=colors.HexColor(MUTED), spaceAfter=2),
        "h2": ParagraphStyle("DnH2", fontName=bold_font, fontSize=13, leading=16, textColor=colors.HexColor(INK), spaceBefore=14, spaceAfter=2),
        "h3": ParagraphStyle("DnH3", fontName=bold_font, fontSize=10.5, leading=14, textColor=colors.HexColor(SOFT), spaceBefore=8, spaceAfter=2),
        "p": ParagraphStyle("DnBody", fontName=body_font, fontSize=10, leading=14.5, textColor=colors.HexColor(INK), spaceAfter=5),
        "li": ParagraphStyle("DnLi", fontName=body_font, fontSize=10, leading=14.5, textColor=colors.HexColor(INK), leftIndent=12, bulletIndent=2, spaceAfter=2),
        "line": ParagraphStyle("DnLine", fontName=body_font, fontSize=9.5, leading=13.5, textColor=colors.HexColor(INK), spaceAfter=3),
    }

    def inline(text: str) -> str:
        return _pdf_inline(text, mono_font, bold_font)

    def rule() -> HRFlowable:
        return HRFlowable(width="100%", thickness=0.6, color=colors.HexColor(LINE), spaceBefore=1, spaceAfter=6)

    story: list = [Paragraph(_escape(session.title), st["title"])]
    meta = [
        f"{ui_message(locale, 'export_start')}: {session.started_at.astimezone().strftime('%d/%m/%Y %H:%M')}",
        f"{ui_message(locale, 'export_duration')}: {_duration(session, segments)}",
    ]
    if participants:
        meta.append(f"{ui_message(locale, 'export_participants')}: {_participants(participants)}")
    story.append(Paragraph(_escape(" · ".join(meta)), st["meta"]))
    story.append(HRFlowable(width="100%", thickness=1.2, color=colors.HexColor(ACCENT), spaceBefore=4, spaceAfter=10))

    blocks = parse_markdown(_summary_markdown(summary, locale))
    if blocks:
        for block in blocks:
            if block.kind == "title":
                continue
            if block.kind == "h2":
                story.append(Paragraph(_escape(block.text), st["h2"]))
                story.append(rule())
            elif block.kind == "h3":
                story.append(Paragraph(_escape(block.text), st["h3"]))
            elif block.kind == "li":
                style = ParagraphStyle(f"li{block.level}", parent=st["li"], leftIndent=12 + 12 * block.level, bulletIndent=2 + 12 * block.level)
                story.append(Paragraph(inline(block.text), style, bulletText="•"))
            else:
                story.append(Paragraph(inline(block.text), st["p"]))
    story.append(Spacer(1, 6))
    story.append(Paragraph(_escape(ui_message(locale, "export_transcript")), st["h2"]))
    story.append(rule())
    if not segments:
        story.append(Paragraph("—", st["p"]))
    for segment in segments:
        speaker = _speaker_name(segment, names)
        who = f'<font name="{bold_font}">{_escape(speaker)}</font> ' if speaker else ""
        stamp = f'<font name="{mono_font}" size="8" color="{MUTED}">{_clock(segment.start_ms)}</font>  '
        story.append(Paragraph(f"{stamp}{who}{_escape(segment.text)}", st["line"]))

    buffer = BytesIO()
    footer_text = f"DroidNote · {session.title}"

    def on_page(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(body_font, 8)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(18 * mm, 10 * mm, footer_text[:110])
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, str(doc.page))
        canvas.restoreState()

    document = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=18 * mm, title=session.title
    )
    document.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buffer.getvalue()


def _pdf_inline(text: str, mono: str, bold: str) -> str:
    """Marcas `mm:ss` e (mm:ss) em mono cinza; **negrito** e `código` do Markdown."""
    out: list[str] = []
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            out.append(f'<font name="{bold}">{_escape(piece[2:-2])}</font>')
        elif piece.startswith("`") and piece.endswith("`"):
            out.append(f'<font name="{mono}" size="8" color="{MUTED}">{_escape(piece[1:-1])}</font>')
        else:
            out.append(_STAMP.sub(lambda m: f'<font name="{mono}" size="8" color="{MUTED}">{m.group(1) or m.group(2)}</font>', _escape(piece)))
    return "".join(out)


# ---------------------------------------------------------------------- DOCX


def export_docx(
    session: Session,
    segments: list[TranscriptSegment],
    summary: Summary | None,
    people: list[Person],
    locale: str | None = None,
    participants: list[Person] | None = None,
) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    accent, ink, soft, muted = RGBColor(0xFF, 0x44, 0x00), RGBColor(0x0C, 0x0C, 0x0E), RGBColor(0x34, 0x34, 0x3A), RGBColor(0x6B, 0x6B, 0x73)
    names = {item.id: item.name for item in people}
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")
    for section in document.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(50)

    def bottom_border(paragraph, color: str, size: int) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(size))
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), color)
        border.append(bottom)
        p_pr.append(border)

    def runs(paragraph, text: str, *, size: float = 10.5, color=ink) -> None:
        """**negrito**, `código` e marcas mm:ss como trechos formatados."""
        for piece in _INLINE.split(text):
            if not piece:
                continue
            if piece.startswith("**") and piece.endswith("**"):
                run = paragraph.add_run(piece[2:-2])
                run.bold = True
                run.font.size = Pt(size)
                run.font.color.rgb = color
            elif piece.startswith("`") and piece.endswith("`"):
                stamp_run(paragraph, piece[1:-1])
            else:
                last = 0
                for match in _STAMP.finditer(piece):
                    if match.start() > last:
                        run = paragraph.add_run(piece[last : match.start()])
                        run.font.size = Pt(size)
                        run.font.color.rgb = color
                    stamp_run(paragraph, match.group(1) or match.group(2))
                    last = match.end()
                if last < len(piece):
                    run = paragraph.add_run(piece[last:])
                    run.font.size = Pt(size)
                    run.font.color.rgb = color

    def stamp_run(paragraph, stamp: str) -> None:
        run = paragraph.add_run(f" {stamp} ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = muted
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")

    title = document.add_paragraph()
    run = title.add_run(session.title)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = accent
    meta = document.add_paragraph()
    parts = [
        f"{ui_message(locale, 'export_start')}: {session.started_at.astimezone().strftime('%d/%m/%Y %H:%M')}",
        f"{ui_message(locale, 'export_duration')}: {_duration(session, segments)}",
    ]
    if participants:
        parts.append(f"{ui_message(locale, 'export_participants')}: {_participants(participants)}")
    meta_run = meta.add_run(" · ".join(parts))
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = muted
    bottom_border(meta, "FF4400", 12)

    def heading(text: str, level: int) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(12 if level == 2 else 8)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(13 if level == 2 else 10.5)
        run.font.color.rgb = ink if level == 2 else soft
        if level == 2:
            bottom_border(paragraph, "D9D9DE", 6)

    for block in parse_markdown(_summary_markdown(summary, locale)):
        if block.kind == "title":
            continue
        if block.kind in {"h2", "h3"}:
            heading(block.text, 2 if block.kind == "h2" else 3)
        elif block.kind == "li":
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Pt(18 + 14 * block.level)
            paragraph.paragraph_format.space_after = Pt(2)
            runs(paragraph, block.text)
        else:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(5)
            runs(paragraph, block.text)

    heading(ui_message(locale, "export_transcript"), 2)
    if not segments:
        document.add_paragraph("—")
    for segment in segments:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        stamp_run(paragraph, _clock(segment.start_ms))
        speaker = _speaker_name(segment, names)
        if speaker:
            run = paragraph.add_run(f"{speaker}  ")
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = soft
        run = paragraph.add_run(segment.text)
        run.font.size = Pt(10)
        run.font.color.rgb = ink
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run(f"DroidNote · {session.title}")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = muted

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------- comum


def _escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _register_pdf_fonts() -> tuple[str, str, str]:
    """(corpo, negrito, mono): fontes do Windows quando existem; senão as padrão do reportlab."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    registered = set(pdfmetrics.getRegisteredFontNames())
    windir = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    wanted = {
        "DnBody": [windir / "segoeui.ttf", windir / "arial.ttf", Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")],
        "DnBold": [windir / "segoeuib.ttf", windir / "arialbd.ttf", Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")],
        "DnMono": [windir / "consola.ttf", windir / "cour.ttf", Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf")],
    }
    fallback = {"DnBody": "Helvetica", "DnBold": "Helvetica-Bold", "DnMono": "Courier"}
    out: list[str] = []
    for name, candidates in wanted.items():
        if name in registered:
            out.append(name)
            continue
        for path in candidates:
            if path.is_file():
                try:
                    pdfmetrics.registerFont(TTFont(name, str(path)))
                    out.append(name)
                    break
                except Exception:
                    continue
        else:
            out.append(fallback[name])
    return out[0], out[1], out[2]
