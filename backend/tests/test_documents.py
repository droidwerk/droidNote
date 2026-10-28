from __future__ import annotations

import io

import pytest
from httpx import AsyncClient

from app.application.chat import ChatService
from app.application.documents import DocumentService
from app.infrastructure.documents.readers import DocumentError, read_document
from app.infrastructure.persistence.sqlite import SqliteStore
from tests.conftest import auth


def _pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for text in pages:
        pdf.drawString(72, 760, text)
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def _docx(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Precos"
    sheet.append(["Plano", "Valor mensal"])
    sheet.append(["Basico", 49])
    sheet.append(["Pro", 129])
    other = book.create_sheet("Prazos")
    other.append(["Entrega", "Data"])
    other.append(["Piloto", "15/11"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def _store(tmp_path) -> SqliteStore:
    store = SqliteStore(tmp_path / "db.sqlite")
    await store.initialize()
    return store


def test_readers_handle_pdf_word_and_excel(tmp_path) -> None:
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(_pdf(["Clausula de multa: 10 por cento", "Vigencia de doze meses"]))
    pieces, pages = read_document(pdf, "contrato.pdf")
    assert pages == 2
    assert [piece.page for piece in pieces] == [1, 2]
    assert "multa" in pieces[0].text

    word = tmp_path / "b.docx"
    word.write_bytes(_docx(["Pauta da reuniao", "Revisar o orcamento de marketing"]))
    pieces, _ = read_document(word, "pauta.docx")
    assert "orcamento de marketing" in pieces[0].text

    excel = tmp_path / "c.xlsx"
    excel.write_bytes(_xlsx())
    pieces, pages = read_document(excel, "precos.xlsx")
    assert pages == 2
    joined = "\n".join(piece.text for piece in pieces)
    assert "Plano: Pro; Valor mensal: 129" in joined
    assert "Aba: Prazos" in joined

    with pytest.raises(DocumentError, match="unsupported"):
        read_document(excel, "planilha.xls")
    broken = tmp_path / "d.pdf"
    broken.write_bytes(b"not a pdf")
    with pytest.raises(DocumentError, match="unreadable"):
        read_document(broken, "quebrado.pdf")


class _CaptureLlm:
    def __init__(self) -> None:
        self.prompt = ""

    async def generate_text(self, prompt: str, *, system: str | None = None) -> str:
        self.prompt = prompt
        return "ok"

    async def generate_text_stream(self, prompt: str, on_token, *, system: str | None = None) -> str:
        self.prompt = prompt
        on_token("ok")
        return "ok"


async def test_chat_answers_from_attached_documents(tmp_path) -> None:
    store = await _store(tmp_path)
    documents = DocumentService(store, tmp_path / "tmp")
    contract = await documents.ingest("contrato.pdf", _pdf(["Clausula de multa: 10 por cento", "Vigencia de doze meses"]))
    prices = await documents.ingest("precos.xlsx", _xlsx())
    assert (contract.pages, contract.chunks) == (2, 2)
    assert list((tmp_path / "tmp").iterdir()) == []  # o temporário some

    llm = _CaptureLlm()
    chat, _, citations = await ChatService(llm, store).ask(
        "Qual a multa do contrato?", document_ids=[contract.id, prices.id]
    )
    assert "[document] contrato.pdf · p. 1" in llm.prompt
    assert "Clausula de multa" in llm.prompt
    assert "precos.xlsx" in llm.prompt
    doc_cites = [item for item in citations if item.document_id]
    assert doc_cites[0].document_name == "contrato.pdf" and doc_cites[0].page == 1
    assert {item.id for item in await store.list_documents(chat_id=chat.id)} == {contract.id, prices.id}

    # Segunda pergunta no mesmo chat: os documentos continuam valendo sem reenviar.
    await ChatService(llm, store).ask("E a vigencia?", chat_id=chat.id)
    assert "Vigencia de doze meses" in llm.prompt
    stored = (await store.list_chat_messages(chat.id))[1]
    assert stored.citations[0].document_name == "contrato.pdf"

    await store.delete_chat(chat.id)
    assert await store.list_documents(chat_id=chat.id) == []


class _OffLlm(_CaptureLlm):
    async def generate_text(self, prompt: str, *, system: str | None = None) -> str:
        raise RuntimeError("motor desligado")

    async def generate_text_stream(self, prompt: str, on_token, *, system: str | None = None) -> str:
        raise RuntimeError("motor desligado")


async def test_failed_first_question_frees_attachments_for_retry(tmp_path) -> None:
    store = await _store(tmp_path)
    prices = await DocumentService(store, tmp_path / "tmp").ingest("precos.xlsx", _xlsx())

    with pytest.raises(RuntimeError):
        await ChatService(_OffLlm(), store).ask("Quanto custa o Pro?", document_ids=[prices.id])
    assert await store.list_chats() == []  # nada de conversa órfã no histórico

    llm = _CaptureLlm()
    chat, _, _ = await ChatService(llm, store).ask("Quanto custa o Pro?", document_ids=[prices.id])
    assert "Valor mensal: 129" in llm.prompt
    assert [item.id for item in await store.list_documents(chat_id=chat.id)] == [prices.id]


async def test_documents_endpoints(client: AsyncClient) -> None:
    created = await client.post(
        "/documents",
        headers=auth(),
        files={"file": ("pauta.docx", _docx(["Pauta: revisar contrato"]), "application/octet-stream")},
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["kind"] == "word" and body["chunks"] == 1

    rejected = await client.post(
        "/documents", headers=auth(), files={"file": ("foto.png", b"\x89PNG", "image/png")}
    )
    assert rejected.status_code == 422 and rejected.json()["detail"] == "unsupported"

    missing_chat = await client.post(
        "/documents", headers=auth(), data={"chat_id": "nope"}, files={"file": ("a.txt", b"oi", "text/plain")}
    )
    assert missing_chat.status_code == 404

    assert (await client.delete(f"/documents/{body['id']}", headers=auth())).status_code == 204
    assert (await client.delete(f"/documents/{body['id']}", headers=auth())).status_code == 404


async def test_live_meeting_question_uses_latest_lines(tmp_path) -> None:
    from datetime import UTC, datetime

    from app.domain.models import Session, TranscriptSegment
    from app.infrastructure.persistence.sqlite import new_id

    store = await _store(tmp_path)
    session = Session(id=new_id(), title="Ao vivo", started_at=datetime.now(tz=UTC), ended_at=None, status="recording")
    await store.create_session(session)
    for index in range(150):
        await store.add_segment(
            TranscriptSegment(
                id=new_id(), session_id=session.id, start_ms=index * 1000, end_ms=index * 1000 + 900,
                text=f"fala numero {index} " + "x" * 120,
            )
        )
    llm = _CaptureLlm()
    await ChatService(llm, store).ask("O que eu perdi?", session_id=session.id)
    assert "fala numero 149" in llm.prompt
    assert "[live meeting]" in llm.prompt
    assert "fala numero 3 " not in llm.prompt
